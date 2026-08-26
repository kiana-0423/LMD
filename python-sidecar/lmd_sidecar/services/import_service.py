from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any


def preview_table_file(file_path: str, preview_rows: int = 20) -> tuple[dict[str, Any], list[str]]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Import file does not exist: {path}")
    if preview_rows < 1 or preview_rows > 1_000:
        raise ValueError("preview_rows must be between 1 and 1000.")

    import pandas as pd

    suffix = path.suffix.lower()
    if suffix == ".csv":
        frame = pd.read_csv(path, nrows=preview_rows)
        sheet_names: list[str] = []
    elif suffix in {".xlsx", ".xlsm"}:
        excel = pd.ExcelFile(path)
        if not excel.sheet_names:
            raise ValueError("The workbook does not contain any worksheets.")
        sheet_names = excel.sheet_names
        frame = pd.read_excel(path, sheet_name=sheet_names[0], nrows=preview_rows)
    else:
        raise ValueError(f"Unsupported import format: {suffix or '(no extension)'}. Use CSV or XLSX.")

    rows = _safe_records(frame)
    return {
        "file_path": str(path),
        "sheet_names": sheet_names,
        "columns": [str(column) for column in frame.columns],
        "rows": rows,
        "preview_rows": rows,
        "preview_count": len(rows),
        "mode": "real",
    }, []


def export_table_rows(file_path: str, output_path: str, chunk_size: int = 500) -> tuple[dict[str, Any], list[str]]:
    """Stream every row to a private JSONL staging file for Rust to ingest."""
    source = Path(file_path)
    destination = Path(output_path)
    if not source.exists():
        raise FileNotFoundError(f"Import file does not exist: {source}")
    if chunk_size < 1 or chunk_size > 10_000:
        raise ValueError("chunk_size must be between 1 and 10000.")
    if not destination.parent.is_dir():
        raise ValueError(f"Import staging directory does not exist: {destination.parent}")

    suffix = source.suffix.lower()
    row_count = 0
    columns: list[str] = []
    with destination.open("w", encoding="utf-8", newline="\n") as stream:
        if suffix == ".csv":
            import pandas as pd

            for frame in pd.read_csv(source, chunksize=chunk_size):
                if not columns:
                    columns = [str(column) for column in frame.columns]
                for record in _safe_records(frame):
                    stream.write(json.dumps(record, ensure_ascii=False, allow_nan=False, default=str) + "\n")
                    row_count += 1
        elif suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                if not workbook.worksheets:
                    raise ValueError("The workbook does not contain any worksheets.")
                worksheet = workbook.worksheets[0]
                values = worksheet.iter_rows(values_only=True)
                header = next(values, None)
                if header is not None:
                    columns = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(header)]
                    for values_row in values:
                        record = {
                            column: _json_safe(value)
                            for column, value in zip(columns, values_row, strict=False)
                        }
                        stream.write(json.dumps(record, ensure_ascii=False, allow_nan=False, default=str) + "\n")
                        row_count += 1
            finally:
                workbook.close()
        else:
            raise ValueError(f"Unsupported import format: {suffix or '(no extension)'}. Use CSV or XLSX.")

    return {
        "file_path": str(source),
        "output_path": str(destination),
        "columns": columns,
        "row_count": row_count,
        "mode": "real",
    }, []


def _safe_records(frame: Any) -> list[dict[str, Any]]:
    return [
        {str(key): _json_safe(value) for key, value in record.items()}
        for record in frame.to_dict("records")
    ]


def _json_safe(value: Any) -> Any:
    if value is None:
        return None
    try:
        import pandas as pd

        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and not math.isfinite(value):
        # json.dumps(allow_nan=False) raises on inf, which would abort the whole import over a
        # single cell. Carry it through as text so the row can still be reviewed.
        return str(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)
